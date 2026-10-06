"""Download the source workbooks and extract the annual series used by the game.

Run locally when a source publishes a new edition. It writes data/data.json, which is committed and
is what the site builds from. Raw workbooks are cached in data/raw/ and are gitignored.

data.json shape:
    {"series": {<key>: {"source": <url>, "description": <str>, "data": [{"year": int, "value": float}]}}}

Rows and columns are located by label text rather than fixed cell positions, so minor layout
changes in future releases shouldn't break the parse. When a new edition is published, update the
URL in SOURCES. GOV.UK's content API lists current attachments, for example
https://www.gov.uk/api/content/government/statistics/energy-trends-section-6-renewables
NESO lists the FES data workbook at https://www.neso.energy/publications/future-energy-scenarios-fes
"""

import datetime
import json
import re
import sys
import urllib.request
from pathlib import Path

import openpyxl

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
RAW_DIR = DATA_DIR / "raw"
OUT_PATH = DATA_DIR / "data.json"

SOURCES = {
    # DESNZ Clean Power 2030 metrics
    "cp2030": "https://assets.publishing.service.gov.uk/media/69cbb7b1b5ac47b0874f3f53/"
    "Clean_Power_2030_Metrics_March_2026.xlsx",
    # Energy Trends section 6: renewables, ET 6.1
    "et61": "https://assets.publishing.service.gov.uk/media/6aba724dfceb6fb3a65012d2/ET_6.1_SEP_26.xlsx",
    # Quarterly Energy Prices: annual domestic electricity bills (QEP 2.2.1) and gas bills (QEP 2.3.1)
    "qep221": "https://assets.publishing.service.gov.uk/media/6a4245907ac6fd9c6a94aae0/table_221__1_.xlsx",
    "qep231": "https://assets.publishing.service.gov.uk/media/6a4245df3413112faed80dbd/table_231__1_.xlsx",
    # Final UK greenhouse gas emissions statistics, 1990 to 2024: data tables
    "ghg": "https://assets.publishing.service.gov.uk/media/6982294819d3abdb495f37ce/"
    "final-greenhouse-gas-emissions-tables-2024.xlsx",
    # DUKES 5.1.2: electricity supply, availability and consumption, 1970 onwards
    "dukes512": "https://assets.publishing.service.gov.uk/media/6a6a36dfbf4a7aca8f8976c8/DUKES_5.1.2.xlsx",
    # NESO Future Energy Scenarios 2025 data workbook (~20 MB)
    "fes": "https://www.neso.energy/document/364551/download",
}


def download(key: str, url: str) -> Path:
    path = RAW_DIR / f"{key}.xlsx"
    # NESO's server rejects requests without a browser-like user agent.
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (clean-power-percent-game)"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        path.write_bytes(resp.read())
    return path


def open_sheet(path: Path, name: str):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if name not in wb.sheetnames:
        raise LookupError(f"Sheet '{name}' not in {path.name}; sheets are {wb.sheetnames}")
    return wb[name]


def as_year(v) -> int | None:
    # Year cells come as ints, dates (FES) or strings, sometimes with revision/provisional markers ("2025 p").
    if isinstance(v, datetime.datetime):
        return v.year
    m = re.fullmatch(r"\s*((?:19|20)\d\d)\s*[a-z\[\]\s]*", str(v)) if v is not None else None
    return int(m.group(1)) if m else None


def as_number(v) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None  # "[x]", blanks and other markers


def row_series(ws, label_prefix: str, scale: float = 1) -> list[dict]:
    """Read a series laid out as a row, with years across a header row somewhere above it.

    The most recent row holding at least three year cells is taken as the header, which handles
    sheets that stack several tables vertically, each with its own header.
    """
    years = None
    for row in ws.iter_rows(values_only=True):
        row_years = [as_year(v) for v in row[1:]]
        if sum(y is not None for y in row_years) >= 3:
            years = row_years
            continue
        label = str(row[0]).strip() if row[0] is not None else ""
        if years and label.startswith(label_prefix):
            return [
                {"year": y, "value": n * scale}
                for y, v in zip(years, row[1:])
                if y is not None and (n := as_number(v)) is not None
            ]
    raise LookupError(f"Row starting '{label_prefix}' not found in sheet '{ws.title}'")


def column_series(ws, column_header: str) -> list[dict]:
    """Read a series laid out as a column, with years down the first column."""
    col = None
    out = []
    for row in ws.iter_rows(values_only=True):
        if col is None:
            headers = [str(v).strip() if v is not None else "" for v in row]
            if column_header in headers:
                col = headers.index(column_header)
            continue
        year, n = as_year(row[0]), as_number(row[col]) if col < len(row) else None
        if year is not None and n is not None:
            out.append({"year": year, "value": n})
    if col is None:
        raise LookupError(f"Column '{column_header}' not found in sheet '{ws.title}'")
    return out


def block_series(ws, header_label: str, row_label: str, scale: float = 1) -> list[dict]:
    """Read a row from a FES chart-data block: a header cell followed by years, then labelled rows below.

    Takes the first block whose header cell starts with `header_label`, and the first row in it labelled
    `row_label` in the same column.
    """
    col = years = None
    for row in ws.iter_rows(values_only=True):
        cells = [str(v).strip() if v is not None else "" for v in row]
        if col is None:
            col = next((i for i, c in enumerate(cells) if c.startswith(header_label)), None)
            if col is not None:
                years = [as_year(v) for v in row[col + 1 :]]
        elif col < len(cells) and cells[col] == row_label:
            return [
                {"year": y, "value": n * scale}
                for y, v in zip(years, row[col + 1 :])
                if y is not None and (n := as_number(v)) is not None
            ]
    raise LookupError(f"Block '{header_label}' / row '{row_label}' not found in sheet '{ws.title}'")


def fes_table_series(ws, item: str, pathway: str, scale: float = 1) -> list[dict]:
    """Read one series from a FES data table, which has 'Data item' and 'Pathway' columns then years."""
    header = None
    for row in ws.iter_rows(values_only=True):
        cells = [str(v).strip() if v is not None else "" for v in row]
        if header is None:
            if "Data item" in cells:
                header = row
                item_col, pathway_col = cells.index("Data item"), cells.index("Pathway")
        elif cells[item_col] == item and cells[pathway_col] == pathway:
            return [
                {"year": y, "value": n * scale}
                for h, v in zip(header, row)
                if (y := as_year(h)) is not None and (n := as_number(v)) is not None
            ]
    raise LookupError(f"'{item}' / '{pathway}' not found in sheet '{ws.title}'")


def since(series: list[dict], first_year: int) -> list[dict]:
    return [d for d in series if d["year"] >= first_year]


def splice(history: list[dict], projection: list[dict]) -> list[dict]:
    """History, then the projection for the years after history ends."""
    last = max(d["year"] for d in history)
    return history + [d for d in projection if d["year"] > last]


def add_series(*series: list[dict]) -> list[dict]:
    """Sum series year by year, keeping only years present in all of them."""
    maps = [{d["year"]: d["value"] for d in s} for s in series]
    common = sorted(set.intersection(*(set(m) for m in maps)))
    return [{"year": y, "value": sum(m[y] for m in maps)} for y in common]


def main() -> int:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    raw = {key: download(key, url) for key, url in SOURCES.items()}

    cp = open_sheet(raw["cp2030"], "Clean Power 2030 Metrics")
    et = open_sheet(raw["et61"], "Annual")
    elec_bill = column_series(open_sheet(raw["qep221"], "2.2.1"), "Overall: All Tariffs (pounds)")
    gas_bill = column_series(open_sheet(raw["qep231"], "2.3.1"), "Overall: All Tariffs (pounds)")
    ghg = open_sheet(raw["ghg"], "1.2")
    dukes = open_sheet(raw["dukes512"], "5.1.2")
    fes_power = open_sheet(raw["fes"], "F.10")
    power_header = "Carbon dioxide equivalent (Mt)"

    series = {
        "clean_generation_share_pct": (
            "cp2030",
            "Clean (nuclear + renewables) share of GB generation, excl. EfW and CHP, %",
            row_series(cp, "Clean share of Great Britain's generation", 100),
        ),
        "clean_demand_share_pct": (
            "cp2030",
            "Share of qualifying GB demand met by clean sources, %",
            row_series(cp, "Share of qualifying GB demand met by clean sources", 100),
        ),
        "emissions_intensity_gco2_kwh": (
            "cp2030",
            "Estimated GB power sector emissions intensity, gCO2/kWh",
            row_series(cp, "Estimated power sector emissions intensity"),
        ),
        "renewables_share_pct": (
            "et61",
            "Renewables share of UK electricity generated (ET 6.1, 'All renewables'), %",
            row_series(et, "All renewables"),
        ),
        "dual_fuel_bill_gbp": (
            "qep221",
            "Average annual UK domestic electricity bill (3,400 kWh) + gas bill (11,200 kWh), "
            "all payment methods and tariffs, cash terms, GBP. Gas from QEP 2.3.1.",
            add_series(elec_bill, gas_bill),
        ),
        "electricity_supply_emissions_mtco2e": (
            "ghg",
            "UK territorial GHG emissions from the electricity supply sector (Table 1.2), MtCO2e",
            row_series(ghg, "Electricity supply total"),
        ),
        "electricity_demand_twh": (
            "dukes512",
            "UK electricity available (supplied + net imports, incl. losses) (DUKES 5.1.2), TWh, 1990 onwards",
            since(column_series(dukes, "Electricity available"), 1990),
        ),
        "fes_ht_demand_twh": (
            "fes",
            "NESO FES 2025 Holistic Transition: GB system demand, total (table ED1), fiscal years, TWh",
            fes_table_series(open_sheet(raw["fes"], "ED1"), "GBFES System Demand: Total", "Holistic Transition", 1e-3),
        ),
        "fes_ht_power_emissions_mtco2e": (
            "fes",
            "NESO FES 2025 power generation emissions (figure F.10): historical, then Holistic Transition, "
            "MtCO2e, 2010 onwards",
            since(
                splice(
                    block_series(fes_power, power_header, "Historical"),
                    block_series(fes_power, power_header, "Holistic Transition"),
                ),
                2010,
            ),
        ),
        "fes_ht_electric_cars_m": (
            "fes",
            "NESO FES 2025 Holistic Transition: battery electric cars on the road (figure F.39), millions",
            block_series(open_sheet(raw["fes"], "F.39"), "#No. of vehicles", "Holistic Transition", 1e-6),
        ),
    }

    out = {
        "series": {
            key: {
                "source": SOURCES[src],
                "description": desc,
                "data": [{"year": d["year"], "value": round(d["value"], 2)} for d in sorted(data, key=lambda d: d["year"])],
            }
            for key, (src, desc, data) in series.items()
        }
    }
    OUT_PATH.write_text(json.dumps(out, indent=2) + "\n", encoding="utf-8")

    for key, s in out["series"].items():
        d = s["data"]
        print(f"{key}: {d[0]['year']}-{d[-1]['year']} ({len(d)} pts), last = {d[-1]['value']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
